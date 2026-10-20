"""Modèles Word modifiables (textes que la Direction voudra adapter)."""
import pypdfium2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from brand import doc, logo
from config import ENTREPRISE as E, DEVISE
from render import BUILD, ROOT, to_pdf

ROUGE, ROUGE_F, GRIS = RGBColor(0xC8, 0x10, 0x2E), RGBColor(0x7E, 0x0A, 0x1E), RGBColor(0x6B, 0x5D, 0x58)
POINTS = "…" * 30


def logo_png():
    html = doc("logo", f'<div style="width:40mm;height:40mm">{logo()}</div>',
               "@page{size:40mm 40mm;margin:0} .logo{width:40mm;height:40mm}", "40mm 40mm")
    pdf = to_pdf(html, BUILD / "logo.pdf")
    png = BUILD / "logo.png"
    pypdfium2.PdfDocument(str(pdf))[0].render(scale=6).to_pil().save(png)
    return png


def ombrer(cell, hex_):
    tc = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_)
    tc.append(shd)


def nouveau(titre, png):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(14), Mm(16)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)

    t = d.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    a, b = t.rows[0].cells
    a.width, b.width = Mm(28), Mm(146)
    a.paragraphs[0].add_run().add_picture(str(png), width=Mm(22))
    p = b.paragraphs[0]
    r = p.add_run("La Tendresse de ")
    r.italic, r.bold, r.font.size = True, True, Pt(16)
    r = p.add_run("KETCHUP")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(20), ROUGE
    q = b.add_paragraph()
    rue = E["rue"] or "rue ……………"
    r = q.add_run(f"{E['type']} · {E['quartier']}, {rue} — {E['ville']}, {E['pays']}\n"
                  f"RCCM : {E['rccm'] or '……………'} · NIU : {E['niu'] or '……………'} · Tél. : {E['telephone'] or '……………'}")
    r.font.size, r.font.color.rgb = Pt(8), GRIS

    h = d.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_before = Pt(14)
    r = h.add_run(titre.upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), ROUGE_F

    pied = sec.footer.paragraphs[0]
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pied.add_run(E["signature"])
    r.italic, r.font.size, r.font.color.rgb = True, Pt(9), ROUGE_F
    return d


def para(d, texte, bold=False, size=11, after=6, align=None):
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.35
    if align:
        p.alignment = align
    r = p.add_run(texte)
    r.bold, r.font.size = bold, Pt(size)
    return p


def sous_titre(d, texte):
    p = d.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(10), Pt(4)
    r = p.add_run(texte)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), ROUGE


def tableau(d, entetes, lignes=0, contenu=None, largeurs=None):
    t = d.add_table(rows=1, cols=len(entetes))
    t.style = "Table Grid"
    for i, h in enumerate(entetes):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), RGBColor(0xFF, 0xFF, 0xFF)
        ombrer(c, "7E0A1E")
    for ligne in (contenu or [[""] * len(entetes)] * lignes):
        cells = t.add_row().cells
        for i, v in enumerate(ligne):
            cells[i].text = v
    if largeurs:
        for row in t.rows:
            for c, w in zip(row.cells, largeurs):
                c.width = Mm(w)
    return t


def signatures(d, noms):
    d.add_paragraph()
    t = d.add_table(rows=2, cols=len(noms))
    for i, n in enumerate(noms):
        t.rows[0].cells[i].text = n
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        t.rows[1].cells[i].text = "\n\n\n"


def reconnaissance_dette(png):
    d = nouveau("Reconnaissance de dette", png)
    for l in [f"Je soussigné(e) {POINTS}{'…' * 25}",
              f"né(e) le {'…' * 15} à {POINTS}",
              f"pièce d'identité (CNI / passeport / autre) n° {POINTS}",
              f"demeurant à {POINTS} téléphone {'…' * 18}",
              f"reconnais devoir au {E['type']} {E['nom']}, sis {E['quartier']}, {E['ville']},",
              f"la somme de (en lettres) {POINTS}{'…' * 10}",
              f"soit (en chiffres) {'…' * 20} {DEVISE},",
              f"au titre de : ☐ consommations   ☐ prêt d'argent   ☐ autre : {'…' * 20}",
              f"Je m'engage à rembourser cette somme au plus tard le {'…' * 15}, selon l'échéancier suivant :"]:
        para(d, l)
    tableau(d, ["Échéance n°", "Date prévue", f"Montant {DEVISE}", "Date payée", "Signature créancier"], 4)
    para(d, "")
    para(d, "En cas de non-paiement à l'échéance, le créancier pourra engager toute procédure de recouvrement prévue par la loi.")
    para(d, f"Fait à {E['ville']}, le {'…' * 15}, en deux exemplaires originaux.")
    signatures(d, ["Le débiteur\n« Lu et approuvé, bon pour la somme de … " + DEVISE + " »",
                   "Pour le créancier\nLa Direction", "Témoin\n(nom et signature)"])
    para(d, "Joindre une photocopie de la pièce d'identité du débiteur.", size=9)
    return d


def reglement_clients(png):
    d = nouveau("Règlement intérieur — Clientèle", png)
    regles = [
        "Toute consommation est payée à la commande. La maison ne fait pas de crédit.",
        "La maison ne sert pas de boissons alcoolisées aux mineurs de moins de 18 ans. Une pièce d'identité peut être demandée.",
        "Le respect du personnel et des autres clients est exigé en toutes circonstances.",
        "Les bagarres, insultes, menaces et comportements dangereux entraînent l'exclusion immédiate.",
        "Aucune boisson ni nourriture extérieure n'est acceptée.",
        "La casse de verres, bouteilles ou mobilier est à la charge du client responsable.",
        "Les armes et produits illicites sont strictement interdits dans l'établissement.",
        "La maison décline toute responsabilité en cas de perte ou de vol d'objets personnels.",
        "Un reçu est remis pour chaque paiement sur simple demande.",
        "Toute réclamation est adressée calmement à la Direction.",
    ]
    for i, r in enumerate(regles, 1):
        para(d, f"Article {i}. {r}")
    para(d, "")
    para(d, "La Direction", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    return d


def reglement_personnel(png):
    d = nouveau("Règles de travail du personnel", png)
    sections = {
        "1. Présence et ponctualité": [
            "Chaque employé respecte les horaires du planning hebdomadaire et signe la fiche de présence.",
            "Toute absence est signalée à l'avance au gérant ; une absence non justifiée n'est pas payée.",
            "Trois retards non justifiés dans le mois font l'objet d'un avertissement."],
        "2. Caisse et argent": [
            "Seule la personne désignée tient la caisse pendant le service.",
            "Toute sortie d'argent fait l'objet d'une pièce de caisse signée, avec justificatif.",
            "Aucun crédit n'est accordé à un client sans accord écrit de la Direction.",
            "La caisse est comptée et clôturée chaque soir ; tout manquant non expliqué peut être retenu sur salaire."],
        "3. Stocks et consommations": [
            "Il est interdit de consommer ou d'offrir une boisson sans bon de sortie interne visé.",
            "Toute casse est déclarée immédiatement dans le registre des casses et pertes.",
            "Les livraisons sont contrôlées avec un bon de réception avant toute signature."],
        "4. Comportement et service": [
            "Tenue propre et correcte exigée. Accueil souriant et courtois de chaque client.",
            "Aucune consommation d'alcool par le personnel pendant le service.",
            "Les conflits avec un client sont remontés au gérant, sans répondre à la provocation."],
        "5. Avances et salaires": [
            "Une seule avance sur salaire en cours par employé, dans la limite de la moitié du salaire.",
            "Le salaire est payé contre signature de la fiche de paiement."],
    }
    for t, items in sections.items():
        sous_titre(d, t)
        for x in items:
            para(d, "•  " + x, after=3)
    para(d, "")
    para(d, f"Lu et approuvé — Nom de l'employé(e) : {POINTS}     Signature : {'…' * 15}")
    para(d, f"Fait à {E['ville']}, le {'…' * 15}.          La Direction", bold=True)
    return d


def facture(png):
    d = nouveau("Facture N° ……………", png)
    para(d, f"Date : {'…' * 15}          Client : {POINTS}")
    para(d, f"Adresse / Tél. client : {POINTS}          NIU client : {'…' * 15}")
    tableau(d, ["N°", "Désignation", "Qté", "Prix unitaire", f"Montant {DEVISE}"], 10, largeurs=[12, 80, 18, 30, 34])
    para(d, "")
    tableau(d, ["", f"Montant {DEVISE}"], contenu=[["Sous-total", ""], ["Remise", ""], ["TOTAL À PAYER", ""]])
    para(d, "")
    para(d, f"Arrêtée la présente facture à la somme de : {POINTS}{'…' * 10} {DEVISE}.")
    para(d, f"Mode de règlement : ☐ Espèces  ☐ MTN MoMo  ☐ Airtel Money  ☐ Virement")
    signatures(d, ["Le Client", "La Direction (signature et cachet)"])
    return d


def relance_credit(png):
    d = nouveau("Lettre de relance — Crédit client", png)
    para(d, f"{E['ville']}, le {'…' * 15}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(d, f"À l'attention de M. / Mme {POINTS}")
    para(d, f"Objet : rappel de paiement — reconnaissance de dette du {'…' * 15}", bold=True)
    para(d, "Madame, Monsieur,")
    para(d, f"Sauf erreur de notre part, la somme de {'…' * 20} {DEVISE}, que vous vous êtes engagé(e) à "
            f"rembourser au plus tard le {'…' * 15}, reste due à ce jour.")
    para(d, f"Nous vous remercions de bien vouloir régulariser cette situation avant le {'…' * 15}, "
            "en espèces à la caisse ou par Mobile Money, contre reçu.")
    para(d, "Si le paiement a été effectué entre-temps, merci de ne pas tenir compte de ce courrier.")
    para(d, "Nous comptons sur votre compréhension et vous prions d'agréer nos salutations distinguées.")
    para(d, "")
    para(d, "La Direction", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    return d


def build(dossier):
    dossier.mkdir(parents=True, exist_ok=True)
    png = logo_png()
    for nom, f in [("Reconnaissance_de_dette", reconnaissance_dette), ("Reglement_interieur_clients", reglement_clients),
                   ("Regles_de_travail_du_personnel", reglement_personnel), ("Facture_modifiable", facture),
                   ("Lettre_de_relance_credit", relance_credit)]:
        f(png).save(dossier / f"{nom}.docx")
    (ROOT / "assets").mkdir(exist_ok=True)
    png.replace(ROOT / "assets" / "logo_la_tendresse_de_ketchup.png")
