"""Classeur Excel de pilotage : caisse, clôtures, stocks, crédits, personnel, tableau de bord."""
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import ENTREPRISE as E, DEVISE
from formulaires import CAT_ENTREE, CAT_SORTIE

ROUGE, ROUGE_F, CREME = "C8102E", "7E0A1E", "FFF4E2"
H_FILL = PatternFill("solid", fgColor=ROUGE_F)
IN_FILL = PatternFill("solid", fgColor="FFFFFF")
CALC_FILL = PatternFill("solid", fgColor=CREME)
H_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
T_FONT = Font(name="Calibri", bold=True, color=ROUGE, size=16)
THIN = Side(style="thin", color="C9B8B2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0 "' + DEVISE + '";-#,##0 "' + DEVISE + '";"-"'
DATE = "dd/mm/yyyy"
N = 600  # lignes préparées par journal

MODES = ["Espèces", "MTN MoMo", "Airtel Money", "Autre"]
MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août",
        "Septembre", "Octobre", "Novembre", "Décembre"]


def titre(ws, texte, sous=""):
    ws["A1"] = f"{E['nom']} — {texte}"
    ws["A1"].font = T_FONT
    ws["A2"] = sous
    ws["A2"].font = Font(italic=True, color="6B5D58", size=9)
    ws.sheet_view.showGridLines = False


def entetes(ws, row, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.fill, cell.font, cell.border = H_FILL, H_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def styler(ws, r0, r1, ncols, calc=(), money=(), dates=()):
    for r in range(r0, r1 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.fill = CALC_FILL if c in calc else IN_FILL
            if c in money:
                cell.number_format = MONEY
            if c in dates:
                cell.number_format = DATE


def liste(ws, dv_range, source):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(dv_range)


def build(dossier):
    dossier.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # ------------------------------------------------ Mode d'emploi
    ws = wb.active
    ws.title = "Mode d'emploi"
    titre(ws, "Classeur de pilotage", f"{E['type']} · {E['quartier']}, {E['ville']}")
    aide = [
        ("Principe", "Les cellules BLANCHES se saisissent. Les cellules CRÈME se calculent toutes seules : ne pas les effacer."),
        ("1. Paramètres", "Saisir le fonds de caisse de départ, la liste du personnel et l'année suivie."),
        ("2. Caisse", "Chaque soir, recopier les pièces de caisse (entrées/sorties) et les ventes du jour (une ligne par mode de paiement)."),
        ("3. Clôtures", "Chaque soir, reporter le solde théorique et les espèces comptées de la fiche de clôture : l'écart se calcule."),
        ("4. Stocks", "Une ligne par article. Chaque semaine, saisir les entrées, sorties et le stock compté de l'inventaire."),
        ("5. Crédits", "Reporter le registre des crédits clients. Le statut (soldé / en retard) se met à jour seul."),
        ("6. Avances", "Reporter les bons d'avance sur salaire."),
        ("7. Tableau de bord", "Rien à saisir : synthèse mensuelle automatique + graphique. À présenter chaque mois à la Direction."),
        ("Règle d'or", "Pas d'argent qui entre ou qui sort sans papier. Un papier = une ligne dans ce classeur."),
    ]
    for i, (a, b) in enumerate(aide, 4):
        ws.cell(row=i, column=1, value=a).font = Font(bold=True, color=ROUGE_F)
        ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 32
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 110

    # ------------------------------------------------ Paramètres
    p = wb.create_sheet("Paramètres")
    titre(p, "Paramètres")
    rows = [("Année suivie", 2026), ("Fonds de caisse de départ", 0), ("Seuil d'alerte écart de caisse", 2000)]
    for i, (a, b) in enumerate(rows, 4):
        p.cell(row=i, column=1, value=a).font = Font(bold=True)
        c = p.cell(row=i, column=2, value=b)
        c.border, c.fill = BORDER, IN_FILL
    p["B5"].number_format = p["B6"].number_format = MONEY
    p.column_dimensions["A"].width = 32
    p.column_dimensions["B"].width = 18
    blocs = [("D", "Catégories entrées", CAT_ENTREE), ("E", "Catégories sorties", CAT_SORTIE),
             ("F", "Modes de paiement", MODES), ("G", "Personnel", [""] * 15)]
    for col, head, items in blocs:
        h = p[f"{col}3"]
        h.value, h.fill, h.font = head, H_FILL, H_FONT
        p.column_dimensions[col].width = 30
        for i, v in enumerate(items, 4):
            c = p[f"{col}{i}"]
            c.value, c.border = v or None, BORDER
    toutes = CAT_ENTREE + CAT_SORTIE
    p["H3"].value, p["H3"].fill, p["H3"].font = "Toutes catégories", H_FILL, H_FONT
    p.column_dimensions["H"].width = 30
    for i, v in enumerate(toutes, 4):
        p[f"H{i}"].value, p[f"H{i}"].border = v, BORDER

    # ------------------------------------------------ Caisse
    c = wb.create_sheet("Caisse")
    titre(c, "Journal de caisse", "Une ligne par pièce de caisse. Le solde se calcule à partir du fonds de départ (Paramètres).")
    cols = ["Date", "N° pièce", "Type", "Catégorie", "Libellé", "Mode", f"Entrée {DEVISE}", f"Sortie {DEVISE}",
            f"Solde {DEVISE}", "Mois"]
    entetes(c, 4, cols, [12, 10, 10, 30, 40, 14, 15, 15, 16, 10])
    for r in range(5, 5 + N):
        c[f"I{r}"] = f'=IF(AND(G{r}="",H{r}=""),"",Paramètres!$B$5+SUM(G$5:G{r})-SUM(H$5:H{r}))'
        c[f"J{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    styler(c, 5, 4 + N, len(cols), calc=(9, 10), money=(7, 8, 9), dates=(1,))
    liste(c, f"C5:C{4 + N}", '"Entrée,Sortie"')
    liste(c, f"D5:D{4 + N}", f"=Paramètres!$H$4:$H${3 + len(CAT_ENTREE + CAT_SORTIE)}")
    liste(c, f"F5:F{4 + N}", "=Paramètres!$F$4:$F$7")
    c.auto_filter.ref = f"A4:J{4 + N}"

    # ------------------------------------------------ Clôtures
    k = wb.create_sheet("Clôtures")
    titre(k, "Clôtures de caisse", "Reporter chaque soir la fiche de clôture.")
    cols = ["Date", "Caissier(ère)", f"Solde théorique {DEVISE}", f"Espèces comptées {DEVISE}", f"Écart {DEVISE}",
            f"MoMo / Airtel {DEVISE}", f"Versé {DEVISE}", "Explication", "Mois"]
    entetes(k, 4, cols, [12, 20, 18, 18, 14, 18, 15, 40, 8])
    for r in range(5, 5 + 400):
        k[f"E{r}"] = f'=IF(D{r}="","",D{r}-C{r})'
        k[f"I{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    styler(k, 5, 404, len(cols), calc=(5, 9), money=(3, 4, 5, 6, 7), dates=(1,))
    k.conditional_formatting.add("E5:E404", FormulaRule(formula=['AND(E5<>"",ABS(E5)>=Paramètres!$B$6)'],
                                                        fill=PatternFill("solid", fgColor="F8C9CF"),
                                                        font=Font(bold=True, color=ROUGE_F)))

    # ------------------------------------------------ Stocks
    s = wb.create_sheet("Stocks")
    titre(s, "Stocks et marges", "Une ligne par article. Mettre à jour à chaque inventaire hebdomadaire.")
    cols = ["Article", "Catégorie", "Unité", f"Prix d'achat {DEVISE}", f"Prix de vente {DEVISE}", "Stock mini",
            "Stock initial", "Entrées", "Ventes", "Casses / offerts", "Stock théorique", "Stock compté", "Écart",
            f"Valeur écart {DEVISE}", f"Valeur stock {DEVISE}", f"Marge / unité {DEVISE}", "Taux de marge", "Alerte"]
    entetes(s, 4, cols, [26, 16, 9, 13, 13, 9, 9, 9, 9, 10, 11, 10, 8, 13, 14, 13, 10, 14])
    exemples = [("Bière 65 cl (à préciser)", "Bières", "Bouteille"), ("Bière 50 cl (à préciser)", "Bières", "Bouteille"),
                ("Boisson sucrée 33 cl", "Sucreries & eaux", "Bouteille"), ("Eau minérale 1,5 L", "Sucreries & eaux", "Bouteille"),
                ("Vin rouge 75 cl", "Vins & spiritueux", "Bouteille"), ("Whisky 70 cl", "Vins & spiritueux", "Bouteille")]
    for i, ex in enumerate(exemples, 5):
        for j, v in enumerate(ex, 1):
            s.cell(row=i, column=j, value=v)
    for r in range(5, 5 + 150):
        s[f"K{r}"] = f'=IF(A{r}="","",N(G{r})+N(H{r})-N(I{r})-N(J{r}))'
        s[f"M{r}"] = f'=IF(L{r}="","",L{r}-K{r})'
        s[f"N{r}"] = f'=IF(M{r}="","",M{r}*N(E{r}))'
        s[f"O{r}"] = f'=IF(A{r}="","",IF(L{r}="",K{r},L{r})*N(D{r}))'
        s[f"P{r}"] = f'=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r})'
        s[f"Q{r}"] = f'=IF(OR(N(E{r})=0,P{r}=""),"",P{r}/E{r})'
        s[f"R{r}"] = f'=IF(OR(A{r}="",F{r}=""),"",IF(IF(L{r}="",K{r},L{r})<=F{r},"À COMMANDER","OK"))'
    styler(s, 5, 154, len(cols), calc=(11, 13, 14, 15, 16, 17, 18), money=(4, 5, 14, 15, 16))
    for r in range(5, 155):
        s[f"Q{r}"].number_format = "0%"
    liste(s, "B5:B154", '"Bières,Sucreries & eaux,Vins & spiritueux,Cuisine,Emballages"')
    s.conditional_formatting.add("R5:R154", CellIsRule(operator="equal", formula=['"À COMMANDER"'],
                                                       fill=PatternFill("solid", fgColor="F8C9CF"),
                                                       font=Font(bold=True, color=ROUGE_F)))
    s.conditional_formatting.add("M5:M154", CellIsRule(operator="lessThan", formula=["0"], font=Font(bold=True, color=ROUGE)))

    # ------------------------------------------------ Crédits
    cr = wb.create_sheet("Crédits clients")
    titre(cr, "Crédits clients", "Crédit exceptionnel, accordé par la Direction uniquement.")
    cols = ["N°", "Date", "Client", "Téléphone", "Détail", f"Montant {DEVISE}", "Échéance", f"Remboursé {DEVISE}",
            f"Reste dû {DEVISE}", "Statut"]
    entetes(cr, 4, cols, [6, 12, 24, 16, 30, 14, 12, 15, 14, 14])
    for r in range(5, 5 + 300):
        cr[f"I{r}"] = f'=IF(F{r}="","",F{r}-N(H{r}))'
        cr[f"J{r}"] = f'=IF(F{r}="","",IF(I{r}<=0,"Soldé",IF(AND(G{r}<>"",TODAY()>G{r}),"EN RETARD","En cours")))'
    styler(cr, 5, 304, len(cols), calc=(9, 10), money=(6, 8, 9), dates=(2, 7))
    cr.conditional_formatting.add("J5:J304", CellIsRule(operator="equal", formula=['"EN RETARD"'],
                                                        fill=PatternFill("solid", fgColor="F8C9CF"),
                                                        font=Font(bold=True, color=ROUGE_F)))

    # ------------------------------------------------ Avances
    av = wb.create_sheet("Avances")
    titre(av, "Avances sur salaire", "Reporter les bons d'avance. Le total du mois est retenu sur la fiche de paiement.")
    cols = ["Date", "N° bon", "Employé(e)", f"Montant {DEVISE}", "Mois de retenue (1-12)", "Retenu ? (Oui/Non)", "Mois"]
    entetes(av, 4, cols, [12, 10, 26, 15, 18, 16, 8])
    for r in range(5, 5 + 300):
        av[f"G{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    styler(av, 5, 304, len(cols), calc=(7,), money=(4,), dates=(1,))
    liste(av, "C5:C304", "=Paramètres!$G$4:$G$18")
    liste(av, "F5:F304", '"Oui,Non"')

    # ------------------------------------------------ Tableau de bord
    t = wb.create_sheet("Tableau de bord")
    titre(t, "Tableau de bord", "Calcul automatique à partir des onglets Caisse, Clôtures, Stocks, Crédits et Avances.")
    cols = ["Mois", f"Ventes {DEVISE}", f"Autres entrées {DEVISE}", f"Dépenses {DEVISE}", f"Résultat {DEVISE}",
            "Résultat / ventes", f"Écarts de caisse {DEVISE}", f"Avances {DEVISE}"] + [f"{x}" for x in CAT_SORTIE]
    entetes(t, 4, cols, [12, 15, 15, 15, 15, 11, 15, 13] + [13] * len(CAT_SORTIE))
    t.freeze_panes = "B5"
    an = "Paramètres!$B$4"
    plage = lambda col: f"Caisse!${col}$5:${col}${4 + N}"
    for m in range(1, 13):
        r = 4 + m
        t[f"A{r}"] = MOIS[m - 1]
        cond = f'{plage("J")},{m},{plage("A")},">="&DATE({an},1,1),{plage("A")},"<="&DATE({an},12,31)'
        t[f"B{r}"] = f'=SUMIFS({plage("G")},{plage("D")},"Ventes du jour",{cond})'
        t[f"C{r}"] = f'=SUMIFS({plage("G")},{cond})-B{r}'
        t[f"D{r}"] = f'=SUMIFS({plage("H")},{cond})'
        t[f"E{r}"] = f"=B{r}+C{r}-D{r}"
        t[f"F{r}"] = f'=IF(B{r}=0,"",E{r}/B{r})'
        t[f"G{r}"] = (f'=SUMIFS(Clôtures!$E$5:$E$404,Clôtures!$I$5:$I$404,{m},'
                      f'Clôtures!$A$5:$A$404,">="&DATE({an},1,1),Clôtures!$A$5:$A$404,"<="&DATE({an},12,31))')
        t[f"H{r}"] = (f'=SUMIFS(Avances!$D$5:$D$304,Avances!$G$5:$G$304,{m},'
                      f'Avances!$A$5:$A$304,">="&DATE({an},1,1),Avances!$A$5:$A$304,"<="&DATE({an},12,31))')
        for j, cat in enumerate(CAT_SORTIE):
            col = get_column_letter(9 + j)
            t[f"{col}{r}"] = f'=SUMIFS({plage("H")},{plage("D")},"{cat}",{cond})'
    t["A17"] = "TOTAL"
    for j in range(2, len(cols) + 1):
        L = get_column_letter(j)
        t[f"{L}17"] = f'=IF(B17=0,"",E17/B17)' if L == "F" else f"=SUM({L}5:{L}16)"
    styler(t, 5, 17, len(cols), calc=tuple(range(1, len(cols) + 1)),
           money=tuple(x for x in range(2, len(cols) + 1) if x != 6))
    for r in range(5, 18):
        t[f"F{r}"].number_format = "0%"
        t[f"A{r}"].font = Font(bold=True)
    for cell in t[17]:
        cell.font = Font(bold=True, color=ROUGE_F)

    t["A20"], t["A21"], t["A22"] = "Crédits clients en cours", "Valeur du stock", "Articles à commander"
    t["B20"] = "=SUM('Crédits clients'!$I$5:$I$304)"
    t["B21"] = "=SUM(Stocks!$O$5:$O$154)"
    t["B22"] = '=COUNTIF(Stocks!$R$5:$R$154,"À COMMANDER")'
    for r in (20, 21, 22):
        t[f"A{r}"].font = Font(bold=True, color=ROUGE_F)
        t[f"B{r}"].fill, t[f"B{r}"].border = CALC_FILL, BORDER
    t["B20"].number_format = t["B21"].number_format = MONEY

    ch = BarChart()
    ch.type = "col"
    ch.title = "Ventes et dépenses par mois"
    ch.y_axis.title = DEVISE
    ch.y_axis.majorGridlines = None
    ch.add_data(Reference(t, min_col=2, min_row=4, max_row=16), titles_from_data=True)
    ch.add_data(Reference(t, min_col=4, min_row=4, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(t, min_col=1, min_row=5, max_row=16))
    ch.series[0].graphicalProperties.solidFill = "1F78B4"
    ch.series[1].graphicalProperties.solidFill = ROUGE
    for se in ch.series:
        se.graphicalProperties.line.noFill = True
    ch.gapWidth, ch.overlap = 60, -10
    ch.legend.position = "b"
    ch.height, ch.width = 8.5, 24
    t.add_chart(ch, "D19")

    wb.move_sheet("Tableau de bord", offset=-(len(wb.sheetnames) - 2))
    out = dossier / "Pilotage_La_Tendresse_de_Ketchup.xlsx"
    wb.save(out)
    return out
